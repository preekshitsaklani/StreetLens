"""StreetLens tests: model identities, reverse valuation, SEC parsing, bridge rules, offline build, Excel parity."""
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from streetlens import bridge, edgar                                   # noqa: E402
from streetlens.cli import build                                        # noqa: E402
from streetlens.models import am_value, bank_value, implied, value_company  # noqa: E402
from streetlens.paths import COMPANIES, SNAP                            # noqa: E402

load = lambda n: json.loads((COMPANIES / f"{n}.json").read_text())
MKT = json.loads((SNAP / "market.json").read_text())
MAC = json.loads((SNAP / "fred.json").read_text())
VAL = load("_valuation")


def test_residual_income_identity():
    """If RoTCE equals Ke every year, no value is created above book: value = tangible book."""
    s = {"rotce1": 0.10, "rotce2": 0.10, "rotce_sust": 0.10, "g": 0.04, "total_payout": 0.6}
    assert abs(bank_value(100.0, 3.0, 0.10, s)["value0"] - 100.0) < 1e-9


def test_dcf_reduces_to_gordon():
    s = {"g1": 0.04, "g2": 0.04, "g3": 0.04, "g4": 0.04, "g5": 0.04, "g": 0.04, "div_payout": 0.4}
    assert abs(am_value(10.0, 0.10, s)["value0"] - 10.0 * 1.04 / 0.06) < 1e-6


def test_reverse_valuation_round_trip():
    c = load("JPM")
    imp = implied(c, MKT, MAC, VAL)
    s = {**c["scenarios"]["base"], "rotce_sust": imp["value"]}
    ke = value_company(c, MKT, MAC, VAL)["ke"]
    assert abs(bank_value(c["latest"]["tbvps"], c["latest"]["dps_run_rate"], ke, s)["tp_raw"] - MKT["companies"]["JPM"]["price"]) < 0.01


def test_sec_history_parsing():
    c = load("JPM")
    h = {x["year"]: x for x in edgar.history([edgar.facts(k) for k in c["ciks"]], c["concepts"])}
    assert h[2025]["eps"] == 20.02 and 0.19 < h[2025]["rotce"] < 0.21


def test_bridge_rules():
    rep = [{"metric": "rotce", "status": "grounded", "quarter": "Q2FY26", "value_pct": 23.0, "evidence": "ROTCE of 23%"}]
    gui = [{"metric": "rotce", "status": "grounded", "quarter": "Q2FY26", "low_pct": 17.0, "high_pct": None, "evidence": "17% through the cycle"}]
    new, log = bridge.apply(load("JPM"), rep, gui)
    assert new["rotce1"] == 0.22 and new["rotce_sust"] == 0.18 and len(log) == 2
    am, _ = bridge.apply(load("BLK"), [{"metric": "organic_base_fee_growth", "status": "grounded", "quarter": "Q2FY26",
                                        "value_pct": 10.0, "evidence": "10% organic base fee growth"}], [])
    assert am["g1"] == 0.145
    same, none = bridge.apply(load("GS"), [], [])
    assert same == load("GS")["scenarios"]["base"] and none == []


def test_offline_build(tmp_path):
    res = build(False, tmp_path)["results"]
    assert {t: (r["rating"], r["tp"]) for t, r in res.items()} == {"JPM": ("HOLD", 320), "GS": ("SELL", 615), "BLK": ("BUY", 1280), "STT": ("HOLD", 155)}
    for f in ("JPM_initiation.md", "GS_initiation.md", "BLK_initiation.md", "US_Financials_sector_note.md",
              "US_Financials_Coverage_Model.xlsx", "valuations.csv", "history.csv"):
        assert (tmp_path / f).exists()


RECALC = Path("/mnt/skills/public/xlsx/scripts/recalc.py")


@pytest.mark.skipif(not (RECALC.exists() and shutil.which("soffice")), reason="LibreOffice recalculation not available")
def test_excel_matches_python(tmp_path):
    from openpyxl import load_workbook
    build(False, tmp_path)
    path = tmp_path / "US_Financials_Coverage_Model.xlsx"
    subprocess.run([sys.executable, str(RECALC), str(path), "90"], check=True, capture_output=True)
    wb = load_workbook(path, data_only=True)
    for t in ("JPM", "GS", "BLK", "STT"):
        py = value_company(load(t), MKT, MAC, VAL)
        assert wb[t]["B36"].value == py["tp"] and wb[t]["B39"].value == py["rating"]


# ---------------------------------------------------------------- real filings, data sources, terminals
REAL = ROOT / "tests" / "fixtures" / "real"


def test_quarter_labels_from_release_dates():
    from streetlens.filings import quarter_label
    assert [quarter_label(d) for d in ("2026-07-14", "2026-04-14", "2026-01-13", "2026-10-13")] == ["Q2FY26", "Q1FY26", "Q4FY25", "Q3FY26"]


def test_grounding_on_a_real_sec_filing():
    """Real JPM 2Q26 earnings release (8-K Exhibit 99.1): true figures pass, an invented one is rejected."""
    from earningslens.extract import RegexExtractor
    from earningslens.grounding import ground
    from earningslens.schema import Extraction, Reported
    text = (REAL / "JPM_Q2FY26_048078.txt").read_text()
    from earningslens.segment import segment
    turns = segment(text)
    assert all(t.role == "management" for t in turns)                  # release mode: no speakers in a filing
    ext = RegexExtractor().run("JPM", "Q2FY26", turns)
    ext.reported.append(Reported(metric="rotce", value=31.0, period="Q2FY26", evidence="an ROTCE of 31%, excluding gains"))
    rep, _ = ground(ext, text, turns)
    status = {(r["metric"], r["value"]): r["status"] for r in rep}
    assert status[("rotce", 23.0)] == "grounded" and status[("cet1_ratio", 14.1)] == "grounded"
    assert status[("rotce", 31.0)] != "grounded"


def test_update_runs_on_real_filings_without_a_key(tmp_path):
    from streetlens.cli import update
    r = update("BLK", 1, "regex", tmp_path, fetch_filings=False, manifest=REAL / "manifest.csv")
    assert r["grounded"] >= 2 and r["rejected"] == 0
    assert [x["driver"] for x in r["log"]] == ["g1"] and r["after"]["tp"] > r["before"]["tp"]
    assert abs(sum(s["usd_per_share"] for s in r["steps"]) - (r["after"]["tp_raw"] - r["before"]["tp_raw"])) < 1e-6
    assert "organic base fee growth" in (tmp_path / "BLK_update.md").read_text()


def test_premium_sources_override_free_ones_field_by_field():
    from streetlens.market import _merge
    out, src = {"companies": {}}, {}
    _merge(out, {"companies": {"JPM": {"price": 330.0, "pb": 2.5, "consensus_tp": 370.0}}}, "yahoo", src)
    _merge(out, {"companies": {"JPM": {"price": 332.4, "pb": None}}}, "bloomberg", src)
    assert out["companies"]["JPM"] == {"price": 332.4, "pb": 2.5, "consensus_tp": 370.0}
    assert src["JPM"] == {"price": "bloomberg", "pb": "yahoo", "consensus_tp": "yahoo"}


def test_beta_from_prices():
    import numpy as np
    from streetlens.market import beta_from_prices
    rng = np.random.default_rng(0)
    m = np.cumsum(rng.normal(0, 0.02, 120))
    idx = [(f"d{i:03d}", float(np.exp(v))) for i, v in enumerate(m)]
    stock = [(d, p ** 1.5) for d, p in idx]                               # log-returns exactly 1.5x the index
    assert beta_from_prices(stock, idx) == 1.5


def test_terminals_switch_off_cleanly_and_formulas_are_written(tmp_path, monkeypatch):
    from streetlens import bloomberg, factset
    from streetlens.terminal_workbook import build as build_terminal
    from openpyxl import load_workbook
    monkeypatch.delenv("FACTSET_API_KEY", raising=False)
    assert factset.available() is False
    try:
        import blpapi  # noqa: F401
    except ImportError:
        assert bloomberg.available() is False
    build_terminal(tmp_path / "t.xlsx", ["JPM"])
    wb = load_workbook(tmp_path / "t.xlsx")
    assert wb["Bloomberg"]["B5"].value == '=BDP("JPM US Equity","PX_LAST")'
    assert wb["FactSet"]["B5"].value == '=FDS("JPM-US","P_PRICE(0)")'
