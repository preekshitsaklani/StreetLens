"""Terminal_Formulas.xlsx: the same inputs as live Excel add-in formulas. Open it on a Bloomberg Terminal
(BDP/BDH) or a FactSet workstation (FDS) and the cells fill in; elsewhere they show #NAME?, which is expected."""
import json

from openpyxl import Workbook
from openpyxl.styles import Font

from .paths import COMPANIES


def build(path, tickers: list[str]) -> None:
    cfg = json.loads((COMPANIES / "_data_fields.json").read_text())
    wb = Workbook()
    bb = wb.active
    bb.title = "Bloomberg"
    b = cfg["bloomberg"]
    bb["A1"] = "Bloomberg Excel add-in: =BDP(security, field) and =BDH(security, field, start, end, options)"
    bb["A1"].font = Font(bold=True)
    bb["A2"] = b["_check"]
    keys = list(b["fields"])
    for j, k in enumerate(["Security"] + keys):
        bb.cell(4, 1 + j, k).font = Font(bold=True)
    for i, t in enumerate(tickers):
        sec = b["security"].format(ticker=t)
        bb.cell(5 + i, 1, sec)
        for j, k in enumerate(keys):
            bb.cell(5 + i, 2 + j, f'=BDP("{sec}","{b["fields"][k]}")')
    r = 6 + len(tickers)
    bb.cell(r, 1, "10-year Treasury"); bb.cell(r, 2, f'=BDP("{b["rf"]}","PX_LAST")')
    bb.cell(r + 1, 1, "Fed funds"); bb.cell(r + 1, 2, f'=BDP("{b["fed_funds"]}","PX_LAST")')
    bb.cell(r + 3, 1, "Weekly closes, 2 years (for beta)").font = Font(bold=True)
    for j, sec in enumerate([b["security"].format(ticker=t) for t in tickers] + [b["index"]]):
        bb.cell(r + 4, 1 + 2 * j, sec)
        bb.cell(r + 5, 1 + 2 * j, f'=BDH("{sec}","PX_LAST","-2CY","","Per=W","Dts=S")')
    fs = wb.create_sheet("FactSet")
    f = cfg["factset"]
    fs["A1"] = "FactSet Excel add-in: =FDS(id, formula)"
    fs["A1"].font = Font(bold=True)
    fs["A2"] = f["_check"]
    keys = list(f["fields"])
    for j, k in enumerate(["Identifier"] + keys):
        fs.cell(4, 1 + j, k).font = Font(bold=True)
    for i, t in enumerate(tickers):
        ident = f["id"].format(ticker=t)
        fs.cell(5 + i, 1, ident)
        for j, k in enumerate(keys):
            fs.cell(5 + i, 2 + j, f'=FDS("{ident}","{f["fields"][k]}")')
    for ws in (bb, fs):
        ws.column_dimensions["A"].width = 22
    wb.save(path)
